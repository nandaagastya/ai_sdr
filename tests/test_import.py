import io
import json
import unittest
from datetime import datetime, timedelta
from unittest.mock import patch

import test_outreach
from app.models import db, Account, Contact, LeadImport, AgentRun, Activity, Task
from app.services.lead_import import parse_file, normalize_rows, suggest_mapping, MAX_BYTES

CSV = b'Company,Email,First name,Job title,Website\nSummit Consulting,jamie@summit.example,Jamie,CEO,summit.example\nSummit Consulting,JAMIE@summit.example,Jamie,CEO,summit.example\nExample Consulting,alex@firm.example,Alex,CEO,firm.example\nBad Row,not-email,Bad,CEO,bad.example\n'


class ImportTests(unittest.TestCase):
    tearDown = test_outreach.OutreachTests.tearDown

    def setUp(self):
        test_outreach.OutreachTests.setUp(self)
        self.client.get('/leads/import')
        with self.client.session_transaction() as session:
            self.token = session['import_csrf']

    def upload(self, content=CSV, name='leads.csv'):
        return self.client.post('/leads/import', data={'csrf_token': self.token,
                                 'file': (io.BytesIO(content), name)}, content_type='multipart/form-data')

    def prepare(self, content=CSV):
        response = self.upload(content)
        self.assertEqual(response.status_code, 303)
        batch = LeadImport.query.order_by(LeadImport.created_at.desc()).first()
        mapping = suggest_mapping(json.loads(batch.data)['headers'])
        response = self.client.post(response.location, data={**mapping, 'csrf_token': self.token})
        self.assertEqual(response.status_code, 303)
        return batch.id, response.location

    @patch('requests.sessions.Session.request', side_effect=AssertionError('Network forbidden'))
    def test_preview_confirm_duplicates_and_repeat_are_safe(self, network):
        original = Account.query.one().to_dict()
        batch_id, url = self.prepare()
        self.assertEqual(Contact.query.count(), 1)
        self.assertEqual(Account.query.count(), 1)
        page = self.client.get(url)
        self.assertIn(b'Import 1 leads', page.data)
        self.assertIn(b'Valid email is required', page.data)
        self.assertIn(b'Email already exists', page.data)
        response = self.client.post(url, data={'csrf_token': self.token})
        self.assertEqual(response.status_code, 303)
        self.assertIn(b'1 leads added', self.client.get(response.location).data)
        self.assertEqual(Contact.query.count(), 2)
        self.assertEqual(Account.query.count(), 2)
        self.assertEqual(db.session.get(Account, original['id']).stage, 'booked')
        self.assertEqual(db.session.get(Account, original['id']).name, original['name'])
        result = json.loads(db.session.get(LeadImport, batch_id).result)
        self.assertEqual(result, {'imported':1,'duplicates':2,'invalid':1,'accounts_created':1})
        self.assertEqual(db.session.get(LeadImport, batch_id).data, '{}')
        self.client.post(url, data={'csrf_token': self.token})
        self.assertEqual(Contact.query.count(), 2)
        self.assertEqual(AgentRun.query.filter_by(agent_name='lead_import').count(), 1)
        self.assertEqual(Activity.query.count(), 0)
        self.assertEqual(Task.query.count(), 0)
        network.assert_not_called()

    def test_existing_company_reused_without_overwrite(self):
        _, url = self.prepare(b'Company,Email,Website\nChanged Name,new@firm.example,firm.example\n')
        self.client.post(url, data={'csrf_token': self.token})
        self.assertEqual(Account.query.count(), 1)
        self.assertEqual(Account.query.one().name, 'Example Consulting')
        self.assertEqual(Account.query.one().stage, 'booked')
        self.assertEqual(Contact.query.count(), 2)

    def test_csrf_session_ownership_and_unreviewed_commit(self):
        self.assertEqual(self.client.post('/leads/import', data={}).status_code, 400)
        response = self.upload()
        batch = LeadImport.query.one()
        other = self.app.test_client()
        self.assertEqual(other.get(response.location).status_code, 404)
        self.assertEqual(other.get(f'/leads/import/{batch.id}').status_code, 404)
        self.assertEqual(self.client.post(response.location, data={}).status_code, 400)
        self.client.post(f'/leads/import/{batch.id}', data={'csrf_token': self.token})
        self.assertEqual(Contact.query.count(), 1)
        batch.created_at = datetime.utcnow() - timedelta(days=2)
        db.session.commit()
        self.assertEqual(self.client.get(response.location).status_code, 410)

    def test_bad_files_and_mapping_do_not_import(self):
        for content, name in [(b'', 'empty.csv'), (b'hello', 'bad.exe'), (b'not zip','bad.xlsx'),
                              (b'Company,Company\nA,B\n', 'duplicate.csv'),
                              (b'Company,Email\nA,a@b.test,extra\n','uneven.csv'),
                              (b'Company,Email\n', 'headers.csv'),
                              (b'x'*(MAX_BYTES+1),'large.csv')]:
            self.assertEqual(self.upload(content, name).status_code, 400)
        response = self.upload()
        self.assertEqual(self.client.post(response.location, data={'csrf_token':self.token,'company':'0','email':'0'}).status_code,400)
        self.assertEqual(self.client.post(response.location, data={'csrf_token':self.token,'company':'0'}).status_code,400)
        self.assertEqual(Contact.query.count(),1)
        self.assertEqual(Account.query.count(),1)

    def test_atomic_rollback_and_recheck_duplicate_at_commit(self):
        batch_id, url = self.prepare()
        self.app.logger.disabled = True
        try:
            with patch.object(db.session, 'commit', side_effect=RuntimeError('disk error')):
                self.assertEqual(self.client.post(url,data={'csrf_token':self.token}).status_code,400)
            self.assertEqual(Contact.query.count(),1)
            self.assertEqual(Account.query.count(),1)
            self.assertEqual(db.session.get(LeadImport,batch_id).status,'reviewed')
            db.session.add(Contact(account_id=Account.query.one().id,email='jamie@summit.example'))
            db.session.commit()
            self.client.post(url,data={'csrf_token':self.token})
            self.assertEqual(Contact.query.count(),2)
            self.assertEqual(json.loads(db.session.get(LeadImport,batch_id).result)['imported'],0)
        finally:
            self.app.logger.disabled = False

    def test_xlsx_docx_and_pdf_parse_real_bytes(self):
        from openpyxl import Workbook
        from docx import Document
        from pypdf import PdfWriter
        from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
        expected = ['Acme','a@acme.test']
        book = Workbook()
        book.active.append(['Company','Email'])
        book.active.append(expected)
        stream = io.BytesIO(); book.save(stream)
        self.assertEqual(parse_file('test.xlsx',stream.getvalue())['rows'],[expected])
        doc = Document(); table = doc.add_table(rows=2,cols=2)
        for i,row in enumerate([['Company','Email'],expected]):
            for j,value in enumerate(row): table.cell(i,j).text=value
        stream = io.BytesIO(); doc.save(stream)
        self.assertEqual(parse_file('test.docx',stream.getvalue())['rows'],[expected])
        writer = PdfWriter(); page = writer.add_blank_page(width=600,height=800)
        font = DictionaryObject({NameObject('/Type'):NameObject('/Font'),NameObject('/Subtype'):NameObject('/Type1'),NameObject('/BaseFont'):NameObject('/Helvetica')})
        page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'):DictionaryObject({NameObject('/F1'):writer._add_object(font)})})
        content = DecodedStreamObject(); content.set_data(b'BT /F1 12 Tf 30 750 Td (Company,Email) Tj 0 -20 Td (Acme,a@acme.test) Tj ET')
        page[NameObject('/Contents')] = writer._add_object(content)
        stream = io.BytesIO(); writer.write(stream)
        self.assertEqual(parse_file('test.pdf',stream.getvalue())['rows'],[expected])
        blank = PdfWriter(); blank.add_blank_page(width=600,height=800)
        stream = io.BytesIO(); blank.write(stream)
        with self.assertRaisesRegex(ValueError,'No readable text'):
            parse_file('scan.pdf',stream.getvalue())

    def test_semicolon_bom_and_validation(self):
        data = parse_file('lead.csv',b'\xef\xbb\xbfCompany;Email;Employees;Website\nAcme;a@acme.test;wrong;javascript://acme.test\n')
        rows = normalize_rows(data,suggest_mapping(data['headers']))
        self.assertEqual(len(rows[0]['errors']),2)
        self.assertIn(b'Company,Email', self.client.get('/leads/import/template.csv').data)


class DashboardTests(unittest.TestCase):
    tearDown = test_outreach.OutreachTests.tearDown

    def setUp(self):
        test_outreach.OutreachTests.setUp(self)

    def test_dashboard_focus_task_completion_and_contact_action(self):
        account = Account.query.one()
        db.session.add(Task(account_id=account.id,description='Prepare for discovery call',owner='human',due_at=datetime.utcnow()-timedelta(days=1)))
        db.session.add(Task(account_id=account.id,description='Already complete',status='done'))
        db.session.commit()
        page = self.client.get('/')
        self.assertEqual(self.client.get('/tasks').status_code, 200)
        self.assertEqual(self.client.get('/leads/discover').status_code, 200)
        self.assertIn(b'Prepare for discovery call',page.data)
        self.assertIn(b'Overdue',page.data)
        self.assertIn(b'Alex',page.data)
        self.assertNotIn(b'Already complete',page.data)
        for removed in (b'Agent stack',b'Data flow',b'Recent agent runs',b'seed.py'):
            self.assertNotIn(removed,page.data)
        task = Task.query.filter_by(status='open').one()
        self.assertEqual(self.client.post(f'/tasks/{task.id}/complete').status_code,400)
        with self.client.session_transaction() as session:
            token=session['dashboard_csrf']
        self.assertEqual(self.client.post(f'/tasks/{task.id}/complete',data={'csrf_token':token}).status_code,303)
        self.assertEqual(db.session.get(Task,task.id).status,'done')
        self.assertIn('You’re caught up'.encode(),self.client.get('/').data)
        outreach = self.client.get(f'/outreach?contact_id={self.contact_id}')
        self.assertIn(f'value="{self.contact_id}" selected'.encode(),outreach.data)
