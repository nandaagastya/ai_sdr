"""Local, bounded lead extraction and duplicate-safe import."""
import csv
import io
import json
import re
import zipfile
from datetime import datetime
from pathlib import Path
from urllib.parse import urlsplit

from app.models import db, Account, Contact, LeadImport, AgentRun

MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 1000
FIELDS = [('company', 'Company *'), ('email', 'Email *'), ('first_name', 'First name'),
          ('last_name', 'Last name'), ('title', 'Job title'), ('domain', 'Company website'),
          ('industry', 'Industry'), ('employee_count', 'Employees'), ('location', 'Location')]
ALIASES = {'company': ['company', 'company name', 'organization', 'account'],
           'email': ['email', 'email address', 'work email'],
           'first_name': ['first name', 'firstname', 'given name'],
           'last_name': ['last name', 'lastname', 'surname'],
           'title': ['title', 'job title', 'position'],
           'domain': ['domain', 'website', 'company website'],
           'industry': ['industry'], 'employee_count': ['employees', 'employee count'],
           'location': ['location', 'country']}


def text_table(text):
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    if not lines:
        raise ValueError('No readable text found. Scans and images need OCR first; upload a CSV or Excel file instead.')
    sample = '\n'.join(lines[:20])
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=',;\t|')
        return list(csv.reader(io.StringIO(text), dialect))
    except csv.Error:
        rows = [re.split(r'\s{2,}', line) for line in lines]
        if len(rows[0]) < 2:
            raise ValueError('No lead table found. Use a table with a header row and one lead per row, or export to CSV.')
        return rows


def parse_file(filename, content):
    extension = Path(filename).suffix.lower()
    if extension not in {'.csv', '.tsv', '.xlsx', '.docx', '.pdf'}:
        raise ValueError('Use CSV, TSV, XLSX, DOCX, or a text-based PDF with a lead table.')
    if not content or len(content) > MAX_BYTES:
        raise ValueError('Choose a non-empty file up to 5 MB.')
    notes = []
    try:
        if extension in {'.xlsx', '.docx'}:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if sum(item.file_size for item in archive.infolist()) > 25 * 1024 * 1024:
                    raise ValueError('Expanded file is too large. Split it into smaller files.')
        if extension in {'.csv', '.tsv'}:
            try:
                text = content.decode('utf-8-sig')
            except UnicodeDecodeError:
                raise ValueError('Save this file as UTF-8 CSV and upload it again.') from None
            rows = text_table(text)
        elif extension == '.xlsx':
            from openpyxl import load_workbook
            book = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
            try:
                sheet = book.worksheets[0]
                if sheet.max_column > 50 or sheet.max_row > MAX_ROWS + 1:
                    raise ValueError('Use at most 1,000 leads and 50 columns per file.')
                rows = list(sheet.iter_rows(values_only=True))
                notes.append(f'Reading first worksheet: {sheet.title}. Additional sheets are not imported.')
            finally:
                book.close()
        elif extension == '.docx':
            from docx import Document
            document = Document(io.BytesIO(content))
            if document.tables:
                rows = [[cell.text for cell in row.cells] for row in document.tables[0].rows]
                notes.append('Reading the first Word table. Additional tables and surrounding text are not imported.')
            else:
                rows = text_table('\n'.join(p.text for p in document.paragraphs))
        else:
            from pypdf import PdfReader
            pdf = PdfReader(io.BytesIO(content))
            if pdf.is_encrypted:
                raise ValueError('Password-protected PDFs are not supported. Upload an unlocked copy.')
            if len(pdf.pages) > 20:
                raise ValueError('Use a PDF with at most 20 pages.')
            texts = []
            for page in pdf.pages:
                stream = page.get_contents()
                if stream is not None and len(stream.get_data()) > 10 * 1024 * 1024:
                    raise ValueError('PDF page is too complex. Export the lead table to CSV.')
                texts.append((page.extract_text(extraction_mode='layout') or '') if stream is not None else '')
            rows = text_table('\n'.join(texts))
            notes.append('PDF layout can shift columns. Carefully check the extracted rows before importing. Scans are not supported.')
    except ValueError:
        raise
    except Exception:
        raise ValueError('Could not read this file. Check the format or export the lead table to CSV.') from None
    rows = [[str(cell).strip() if cell is not None else '' for cell in row] for row in rows
            if any(cell is not None and str(cell).strip() for cell in row)]
    if len(rows) < 2:
        raise ValueError('Include a header row followed by at least one lead.')
    headers = rows[0]
    if len(headers) > 50 or len(rows) - 1 > MAX_ROWS:
        raise ValueError('Use at most 1,000 leads and 50 columns per file.')
    if any(len(cell) > 5000 for row in rows for cell in row):
        raise ValueError('A cell exceeds 5,000 characters. Upload only the lead table.')
    if not all(headers) or len(set(h.casefold() for h in headers)) != len(headers):
        raise ValueError('Every column needs a unique, non-empty header.')
    cleaned = []
    for number, row in enumerate(rows[1:], 2):
        if row == headers:  # Repeated PDF table headers.
            continue
        if len(row) != len(headers):
            raise ValueError(f'Row {number} has a different number of columns. Correct the table or export it to CSV.')
        cleaned.append(row)
    if not cleaned:
        raise ValueError('No lead rows found after the header.')
    return {'headers': headers, 'rows': cleaned, 'notes': notes}


def suggest_mapping(headers):
    normalized = [re.sub(r'[_-]+', ' ', value.casefold()).strip() for value in headers]
    return {field: next((str(i) for i, h in enumerate(normalized) if h in ALIASES[field]), '')
            for field, _ in FIELDS}


def normalize_rows(data, mapping):
    indexes = {}
    for field, _ in FIELDS:
        value = mapping.get(field, '')
        if value == '':
            if field in {'company', 'email'}:
                raise ValueError('Map both Company and Email before continuing.')
            continue
        try:
            index = int(value)
        except (ValueError, TypeError):
            raise ValueError('Choose a column from this file.') from None
        if index < 0 or index >= len(data['headers']) or index in indexes.values():
            raise ValueError('Map each file column to only one field.')
        indexes[field] = index
    result = []
    for number, raw in enumerate(data['rows'], 2):
        row = {field: raw[index].strip() for field, index in indexes.items()}
        errors = []
        for field, limit in [('company',255), ('email',255), ('first_name',120), ('last_name',120),
                             ('title',255), ('domain',255), ('industry',120), ('location',255)]:
            value = row.get(field, '')
            if len(value) > limit or any(ord(c) < 32 for c in value):
                errors.append(f'{field.replace("_", " ")} is too long or contains line breaks')
        if not row.get('company'):
            errors.append('Company is required')
        row['email'] = row.get('email', '').casefold()
        if not re.fullmatch(r'[^\s@]+@[^\s@]+\.[^\s@]+', row['email']):
            errors.append('Valid email is required')
        domain = row.get('domain', '')
        if domain:
            parsed = urlsplit(domain if '://' in domain else 'https://' + domain)
            if parsed.scheme not in {'http', 'https'} or not parsed.hostname or '.' not in parsed.hostname or parsed.username or re.search(r'\s', domain):
                errors.append('Company website is invalid')
            else:
                row['domain'] = parsed.hostname.lower().removeprefix('www.')
        employees = row.get('employee_count', '')
        if employees:
            if not employees.isdigit() or not 1 <= int(employees) <= 100000000:
                errors.append('Employees must be a positive whole number')
            else:
                row['employee_count'] = int(employees)
        else:
            row['employee_count'] = None
        result.append({'row': number, 'lead': row, 'errors': errors})
    return result


def classify(rows):
    existing = {email.casefold() for (email,) in db.session.query(Contact.email).filter(Contact.email.isnot(None))}
    seen = set()
    output = []
    for item in rows:
        lead = item['lead']
        reason, state = '; '.join(item['errors']), 'invalid' if item['errors'] else 'ready'
        email = lead['email']
        if state == 'ready':
            if email in existing or email in seen:
                state, reason = 'duplicate', 'Email already exists or repeats in this file'
            else:
                seen.add(email)
        output.append({**item, 'state': state, 'reason': reason})
    return output


def commit_import(batch):
    if batch.status == 'complete':
        return json.loads(batch.result)
    if batch.status != 'reviewed':
        raise ValueError('Review and map the file before importing.')
    rows = json.loads(batch.data)['leads']
    try:
        claimed = db.session.execute(db.update(LeadImport).where(
            LeadImport.id == batch.id, LeadImport.status == 'reviewed').values(status='importing'),
            execution_options={'synchronize_session': False})
        if claimed.rowcount != 1:
            raise ValueError('This import is already being processed. Reload the page.')
        counts = {'imported': 0, 'duplicates': 0, 'invalid': 0, 'accounts_created': 0}
        for item in classify(rows):
            if item['state'] != 'ready':
                counts['duplicates' if item['state'] == 'duplicate' else 'invalid'] += 1
                continue
            lead = item['lead']
            account = None
            if lead.get('domain'):
                account = Account.query.filter(db.func.lower(Account.domain) == lead['domain']).first()
            if account is None:
                candidates = Account.query.filter(db.func.lower(Account.name) == lead['company'].lower()).all()
                candidates = [a for a in candidates if not lead.get('domain') or not a.domain or a.domain == lead['domain']]
                if len(candidates) > 1:
                    counts['invalid'] += 1
                    continue
                account = candidates[0] if candidates else None
            if account is None:
                account = Account(name=lead['company'], domain=lead.get('domain') or None,
                                  industry=lead.get('industry') or None, employee_count=lead.get('employee_count'),
                                  location=lead.get('location') or None, source='import', stage='new',
                                  icp_fit_score=0, icp_fit_reasons='Imported lead; fit not yet assessed')
                db.session.add(account)
                db.session.flush()
                counts['accounts_created'] += 1
                from app.services.lead_quality import apply_quality
                apply_quality(account)
            db.session.add(Contact(account_id=account.id, email=lead['email'],
                                   first_name=lead.get('first_name') or None, last_name=lead.get('last_name') or None,
                                   title=lead.get('title') or None))
            counts['imported'] += 1
        batch.status = 'complete'
        batch.result = json.dumps(counts)
        # Discard staged lead rows after confirmation; imported contacts are authoritative.
        batch.data = '{}'
        db.session.add(AgentRun(agent_name='lead_import', status='success', records_processed=counts['imported'],
                                detail=f'Lead file import {batch.id}: {json.dumps(counts)}. No outreach sent.',
                                finished_at=datetime.utcnow()))
        db.session.commit()
        return counts
    except Exception:
        db.session.rollback()
        raise
